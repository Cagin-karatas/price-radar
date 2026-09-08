"""CSV ve Excel dışa aktarımı.

Üç veri kümesi:
  - `products` : ürünler, teklif sayısı ve fiyat aralığıyla
  - `offers`   : teklifler, son fiyat ve stok durumuyla
  - `history`  : fiyat geçmişi zaman serisi

Excel biçimlendirmesi Proje 1'deki (exceltool) yaklaşımın aynısı: başlık dolgusu,
donmuş satır, otomatik filtre, hesaplanmış sütun genişliği.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..db.models import Offer, PriceSnapshot, Product, Site

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LINK_FONT = Font(color="0563C1", underline="single")

DATASETS = ("products", "offers", "history")


async def products_frame(session: AsyncSession, min_offers: int = 1) -> pd.DataFrame:
    """Ürünler: teklif sayısı, en düşük/en yüksek fiyat, fark."""
    rows = (
        await session.execute(
            select(
                Product.id,
                Product.title,
                Product.brand,
                Product.category,
                func.count(Offer.id).label("teklif_sayisi"),
                func.min(Offer.last_price).label("en_dusuk"),
                func.max(Offer.last_price).label("en_yuksek"),
                func.min(Offer.last_currency).label("para_birimi"),
            )
            .join(Offer, Offer.product_id == Product.id)
            .group_by(Product.id, Product.title, Product.brand, Product.category)
            .having(func.count(Offer.id) >= min_offers)
            .order_by(func.count(Offer.id).desc(), Product.title)
        )
    ).all()

    frame = pd.DataFrame(
        [
            {
                "urun_id": row.id,
                "urun": row.title,
                "marka": row.brand,
                "kategori": row.category,
                "teklif_sayisi": row.teklif_sayisi,
                "en_dusuk": row.en_dusuk,
                "en_yuksek": row.en_yuksek,
                "para_birimi": row.para_birimi,
            }
            for row in rows
        ]
    )

    if frame.empty:
        return frame

    # Fiyat farkı: siteler arası tasarruf potansiyeli — tablonun asıl değeri
    frame["fark"] = frame["en_yuksek"] - frame["en_dusuk"]
    frame["fark_yuzde"] = (
        (frame["fark"] / frame["en_yuksek"].replace(0, pd.NA) * 100).round(1)
    )
    return frame.sort_values("fark", ascending=False, na_position="last")


async def offers_frame(session: AsyncSession, site: str | None = None) -> pd.DataFrame:
    """Teklifler: hangi üründen, hangi sitede, son fiyat ve stok."""
    query = (
        select(
            Offer.id,
            Product.title.label("urun"),
            Site.slug.label("site"),
            Offer.raw_title,
            Offer.last_price,
            Offer.last_currency,
            Offer.last_availability,
            Offer.last_checked_at,
            Offer.url,
        )
        .join(Product, Product.id == Offer.product_id)
        .join(Site, Site.id == Offer.site_id)
        .order_by(Product.title, Offer.last_price)
    )
    if site:
        query = query.where(Site.slug == site)

    rows = (await session.execute(query)).all()

    return pd.DataFrame(
        [
            {
                "teklif_id": row.id,
                "urun": row.urun,
                "site": row.site,
                "site_basligi": row.raw_title,
                "fiyat": row.last_price,
                "para_birimi": row.last_currency,
                "stok": row.last_availability,
                "son_kontrol": row.last_checked_at,
                "baglanti": row.url,
            }
            for row in rows
        ]
    )


async def history_frame(
    session: AsyncSession, product_id: int | None = None, limit: int = 10_000
) -> pd.DataFrame:
    """Fiyat geçmişi zaman serisi."""
    query = (
        select(
            PriceSnapshot.captured_at,
            Product.id.label("urun_id"),
            Product.title.label("urun"),
            Site.slug.label("site"),
            PriceSnapshot.price,
            PriceSnapshot.currency,
            PriceSnapshot.availability,
            Offer.id.label("teklif_id"),
        )
        .join(Offer, Offer.id == PriceSnapshot.offer_id)
        .join(Product, Product.id == Offer.product_id)
        .join(Site, Site.id == Offer.site_id)
        .order_by(PriceSnapshot.captured_at.desc())
        .limit(limit)
    )
    if product_id:
        query = query.where(Product.id == product_id)

    rows = (await session.execute(query)).all()

    return pd.DataFrame(
        [
            {
                "tarih": row.captured_at,
                "urun_id": row.urun_id,
                "urun": row.urun,
                "site": row.site,
                "fiyat": row.price,
                "para_birimi": row.currency,
                "stok": row.availability,
                "teklif_id": row.teklif_id,
            }
            for row in rows
        ]
    )


def _strip_timezone(frame: pd.DataFrame) -> pd.DataFrame:
    """Excel saat dilimli tarihleri kabul etmiyor; yerelleştirmeyi kaldır."""
    out = frame.copy()
    for column in out.columns:
        series = out[column]
        if pd.api.types.is_datetime64_any_dtype(series) and getattr(series.dt, "tz", None):
            out[column] = series.dt.tz_localize(None)
    return out


def to_csv(frame: pd.DataFrame) -> bytes:
    """UTF-8 BOM ile: Excel Türkçe karakterleri doğru açsın."""
    return frame.to_csv(index=False).encode("utf-8-sig")


def _style(worksheet, frame: pd.DataFrame, link_column: str | None = None) -> None:
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.freeze_panes = "A2"
    if len(frame) and frame.shape[1]:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(frame.shape[1])}{len(frame) + 1}"

    for index, column in enumerate(frame.columns, start=1):
        sample = frame[column].head(200)
        # NA değerler astype(str) sonrası da NA kalabiliyor; len() sarmalanıyor
        lengths = sample.map(lambda v: len(str(v)) if v is not None and v == v else 0)
        width = min(max(len(str(column)), int(lengths.max()) if len(lengths) else 0) + 3, 50)
        worksheet.column_dimensions[get_column_letter(index)].width = max(width, 10)

    if link_column and link_column in frame.columns:
        col_index = list(frame.columns).index(link_column) + 1
        for row in range(2, len(frame) + 2):
            cell = worksheet.cell(row=row, column=col_index)
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.font = LINK_FONT
                cell.value = "Ürüne git"


def to_excel(sheets: dict[str, pd.DataFrame], links: dict[str, str] | None = None) -> bytes:
    """Çok sayfalı, biçimlendirilmiş Excel üretir."""
    links = links or {}
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            safe_name = name[:31]
            clean = _strip_timezone(frame)
            if clean.empty:
                # Boş sayfa yerine açıklayıcı bir satır: dosyayı açan
                # kişi verinin gelmediğini anlasın
                clean = pd.DataFrame({"bilgi": ["Bu veri kümesinde kayıt yok"]})
            clean.to_excel(writer, sheet_name=safe_name, index=False)
            _style(writer.sheets[safe_name], clean, links.get(name))

    return buffer.getvalue()


async def build_export(
    session: AsyncSession,
    dataset: str,
    fmt: str,
    *,
    site: str | None = None,
    product_id: int | None = None,
) -> tuple[bytes, str, str]:
    """(içerik, dosya adı, MIME türü) döndürür."""
    if dataset not in DATASETS:
        raise ValueError(f"Bilinmeyen veri kümesi: {dataset}. Seçenekler: {', '.join(DATASETS)}")
    if fmt not in {"csv", "xlsx"}:
        raise ValueError(f"Bilinmeyen biçim: {fmt}. Seçenekler: csv, xlsx")

    if dataset == "products":
        frame = await products_frame(session)
    elif dataset == "offers":
        frame = await offers_frame(session, site=site)
    else:
        frame = await history_frame(session, product_id=product_id)

    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    filename = f"priceradar-{dataset}-{stamp}.{fmt}"

    if fmt == "csv":
        return to_csv(frame), filename, "text/csv; charset=utf-8"

    links = {"offers": "baglanti"}
    content = to_excel({dataset: frame}, links=links)
    return (
        content,
        filename,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


async def build_workbook(session: AsyncSession) -> tuple[bytes, str, str]:
    """Üç veri kümesini tek Excel dosyasında toplar."""
    sheets = {
        "Urunler": await products_frame(session),
        "Teklifler": await offers_frame(session),
        "Fiyat gecmisi": await history_frame(session),
    }
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    return (
        to_excel(sheets, links={"Teklifler": "baglanti"}),
        f"priceradar-{stamp}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
